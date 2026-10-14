#!/usr/bin/env python3
"""Cria ficheiros de exemplo (docx, xlsx, pptx) numa pasta, para testar a conversão.
Uso:  python tests/criar_amostras.py <pasta>
Requer:  pip install python-docx openpyxl python-pptx
"""
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches as PInches

pasta = Path(sys.argv[1] if len(sys.argv) > 1 else "amostras")
(pasta / "sub").mkdir(parents=True, exist_ok=True)

# Word
d = Document()
d.add_heading("Relatório de teste", 0)
d.add_paragraph("Texto com acentuação: coração, ação. Símbolos: ✓ ★ Ω € ±")
for i in range(40):
    d.add_paragraph(f"Linha {i} do documento Word de teste.")
d.save(pasta / "relatorio.docx")
d2 = Document(); d2.add_paragraph("Documento na subpasta."); d2.save(pasta / "sub" / "interno.docx")

# Excel: duas folhas largas (26 colunas); uma com impressão ajustada a 1 página de largura
def folha_larga(ajustar):
    wb = Workbook(); ws = wb.active; ws.title = "Dados"
    ws.append(["Item"] + [f"Coluna {i}" for i in range(1, 26)])
    for r in range(1, 31):
        ws.append([f"Linha {r}"] + [r * c for c in range(1, 26)])
    if ajustar:
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    return wb
folha_larga(False).save(pasta / "larga_normal.xlsx")
folha_larga(True).save(pasta / "larga_ajustada.xlsx")

# PowerPoint
pr = Presentation()
for t in ("Slide um", "Slide dois", "Slide três"):
    s = pr.slides.add_slide(pr.slide_layouts[5]); s.shapes.title.text = t
    tb = s.shapes.add_textbox(PInches(1), PInches(2), PInches(6), PInches(1))
    tb.text_frame.text = "Conteúdo de teste ✓ €"
pr.save(pasta / "apresentacao.pptx")
print("Amostras criadas em", pasta)
